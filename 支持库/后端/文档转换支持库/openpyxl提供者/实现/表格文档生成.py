"""openpyxl 独立提供者生成实现：内容参数 → XLSX 字节 → 生成产物字典。

内容参数：{"工作表列表": [{"表名"|"name"|"标题", "列"|"columns": [...],
"行"|"rows": [...]}]}，兼容单元格字典（文本/标注/value/name）。稳定
错误码：参数不合法/提供者不可用/生成失败。
"""

from __future__ import annotations

import base64
import hashlib
import io
import time

import 支持库.后端.文档转换支持库.openpyxl提供者.实现.表格文档 as _公共
from 公共契约.基础类型.结果类型 import 结果

媒体类型_xlsx = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _取表名(工作表: dict) -> str:
    名称 = 工作表.get("表名") or 工作表.get("name") or 工作表.get("标题") or "工作表"
    return str(名称)[:31] or "工作表"


def _取单元格文本(值) -> str:
    if isinstance(值, dict):
        for 键 in ("文本", "标注", "value", "name", "名称"):
            if 键 in 值 and 值[键] is not None:
                return str(值[键])
        return str(值)
    return "" if 值 is None else str(值)


def _列名(列单元格) -> str:
    """列可以是文本或 {名称: 列名} 字典。"""
    return _取单元格文本(列单元格)


def _标准化行(行数据, 列: list) -> list:
    if isinstance(行数据, dict):
        if 列:
            return [_取单元格文本(行数据.get(_列名(列单元格))) for 列单元格 in 列]
        return [_取单元格文本(值) for 值 in 行数据.values()]
    if isinstance(行数据, (list, tuple)):
        return [_取单元格文本(值) for 值 in 行数据]
    return [_取单元格文本(行数据)]


def _解析单元格地址(地址: str) -> tuple[int, int]:
    """把 A1 风格地址解析为 (行号, 列号)，非法时抛 ValueError。"""
    匹配 = __import__("re").match(r"([A-Za-z]+)(\d+)", str(地址 or "").strip())
    if not 匹配:
        raise ValueError(f"非法单元格地址: {地址}")
    列索引 = 0
    for 字符 in 匹配.group(1).upper():
        列索引 = 列索引 * 26 + (ord(字符) - 64)
    return int(匹配.group(2)), 列索引


def _应用样式到单元格(单元格, 样式信息: dict) -> None:
    """把样式字典应用到 openpyxl 单元格（与旧版写入_样式 同语义）。"""
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    字体参数: dict = {}
    if 样式信息.get("bold"):
        字体参数["bold"] = True
    if 样式信息.get("italic"):
        字体参数["italic"] = True
    if 样式信息.get("underline"):
        字体参数["underline"] = "single"
    if 样式信息.get("strikethrough"):
        字体参数["strike"] = True
    if 样式信息.get("fontSize"):
        字体参数["size"] = 样式信息["fontSize"]
    if 样式信息.get("fontName"):
        字体参数["name"] = 样式信息["fontName"]
    if 样式信息.get("color"):
        颜色值 = str(样式信息["color"])
        if 颜色值.startswith("#"):
            字体参数["color"] = 颜色值[1:]
    if 字体参数:
        单元格.font = Font(**字体参数)
    if 样式信息.get("fillColor"):
        填充色 = str(样式信息["fillColor"])
        if 填充色.startswith("#"):
            单元格.fill = PatternFill(start_color=填充色[1:], end_color=填充色[1:], fill_type="solid")
    对齐参数: dict = {}
    if 样式信息.get("对齐"):
        水平对齐 = {"左": "left", "居中": "center", "右": "right"}.get(样式信息["对齐"])
        if 水平对齐:
            对齐参数["horizontal"] = 水平对齐
    if 样式信息.get("wrapText"):
        对齐参数["wrap_text"] = True
    if 对齐参数:
        单元格.alignment = Alignment(**对齐参数)
    if 样式信息.get("borderType"):
        边框类型 = 样式信息["borderType"]
        if 边框类型 in ("全部", "outside"):
            单元格.border = Border(
                left=Side(style="thin", color="000000"),
                right=Side(style="thin", color="000000"),
                top=Side(style="thin", color="000000"),
                bottom=Side(style="thin", color="000000"),
            )


def _写入单元格映射(表单, 工作表: dict) -> int:
    """按 A1 地址写入 单元格映射（含公式/样式/合并/列宽/行高）。"""
    from openpyxl.styles import Alignment, Font
    写入格数 = 0
    单元格映射 = 工作表.get("单元格映射") or {}
    样式映射 = 工作表.get("样式映射") or {}
    if 单元格映射:
        for 地址, 值 in 单元格映射.items():
            try:
                行号, 列号 = _解析单元格地址(地址)
            except ValueError:
                continue
            单元格 = 表单.cell(row=行号, column=列号)
            if isinstance(值, str) and "[ASCII:" in 值:
                值 = _解码ascii转义(值)
            单元格.value = 值
            样式信息 = 样式映射.get(地址)
            if isinstance(样式信息, dict) and 样式信息:
                _应用样式到单元格(单元格, 样式信息)
            写入格数 += 1
    for 合并范围 in 工作表.get("合并范围") or []:
        if isinstance(合并范围, str):
            try:
                表单.merge_cells(合并范围)
            except Exception as 错误:
                continue
    for 列字母, 宽度 in (工作表.get("列宽映射") or {}).items():
        try:
            表单.column_dimensions[str(列字母)].width = float(宽度) / 7
        except (KeyError, ValueError, TypeError):
            continue
    for 行号字符串, 高度 in (工作表.get("行高映射") or {}).items():
        try:
            表单.row_dimensions[int(行号字符串)].height = float(高度) / 4
        except (ValueError, TypeError, KeyError):
            continue
    return 写入格数


def _解码ascii转义(值: str) -> str:
    """解码 [ASCII:0xXXXX] 转义序列（旧版 Excel 引擎产物兼容）。"""
    import re
    def 替换(匹配):
        try:
            return chr(int(匹配.group(1), 16))
        except ValueError:
            return 匹配.group(0)
    return re.sub(r"\[ASCII:0x([0-9A-F]{4})\]", 替换, 值)


def _生成XLSX字节(工作表列表: list) -> bytes:
    """按 工作表列表 生成 XLSX 字节；无可写单元格时抛 ValueError。"""
    from openpyxl import Workbook
    工作簿 = Workbook()
    工作簿.remove(工作簿.active)
    写入格数 = 0
    for 工作表 in 工作表列表:
        if not isinstance(工作表, dict):
            continue
        列 = 工作表.get("列") or 工作表.get("columns") or []
        行 = 工作表.get("行") or 工作表.get("rows") or []
        列 = 列 if isinstance(列, list) else []
        行 = 行 if isinstance(行, list) else []
        表单 = 工作簿.create_sheet(title=_取表名(工作表))
        if 列:
            表单.append([_取单元格文本(单元格) for 单元格 in 列])
            写入格数 += len(列)
        for 行数据 in 行:
            行值 = _标准化行(行数据, 列)
            if 行值:
                表单.append(行值)
                写入格数 += len(行值)
        if 工作表.get("单元格映射"):
            写入格数 += _写入单元格映射(表单, 工作表)
    if 写入格数 == 0:
        raise ValueError("工作表列表中没有可写单元格（列/行/单元格映射均为空）")
    缓冲 = io.BytesIO()
    工作簿.save(缓冲)
    return 缓冲.getvalue()


def 生成表格文档(内容参数: dict) -> 结果:
    """按 工作表列表 生成 XLSX，返回生成产物字典（字节b64/媒体类型/摘要）。"""
    if not isinstance(内容参数, dict):
        return _公共._失败("参数不合法", "内容参数必须是字典")
    提供者 = _公共.加载提供者()
    if 提供者["openpyxl"] is None:
        return _公共._失败("提供者不可用", "openpyxl 不可用，无法生成表格文档", 可重试=True)
    工作表列表 = 内容参数.get("工作表列表") or []
    if not isinstance(工作表列表, list) or not 工作表列表:
        return _公共._失败("参数不合法", "工作表列表不能为空，至少需要一个工作表")
    开始时刻 = time.monotonic()
    try:
        字节 = _生成XLSX字节(工作表列表)
    except ValueError as 错误:
        return _公共._失败("参数不合法", str(错误))
    except Exception as 错误:
        return _公共._失败("生成失败", f"XLSX 生成异常: {错误}")
    if not 字节:
        return _公共._失败("生成失败", "XLSX 生成结果为空")
    return 结果.成功结果({
        "格式": "xlsx",
        "字节b64": base64.b64encode(字节).decode("ascii"),
        "媒体类型": 媒体类型_xlsx,
        "摘要": hashlib.sha256(字节).hexdigest(),
        "诊断": [f"XLSX 生成成功，耗时 {time.monotonic() - 开始时刻:.3f} 秒"],
        "提供者版本": dict(提供者["版本"]),
        "字节数": len(字节),
    })
