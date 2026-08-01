"""演示文稿解析原子能力：pptx 原生 + ppt 经文档转换链，返回 结果[通用文档]。安全约束：OOXML 不可信 ZIP 检查；损坏/伪装→文件损坏；超限→超出限制；缺 python-pptx→提供者不可用。"""

from __future__ import annotations

import base64
import hashlib
import shutil
import tempfile
import time
import zipfile
from pathlib import Path
from typing import Any

from 公共契约.基础类型.文档结构 import 文档块, 文档资源, 来源位置, 通用文档, 归一化格式
from 公共契约.基础类型.结果类型 import 结果
from 支持库.后端.文档转换 import 转换办公文件

try:
    import pptx
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    pptx = MSO_SHAPE_TYPE = None

默认最大幻灯片数 = 200
默认最大字节数 = 200 * 1024 * 1024
默认超时秒 = 60
最大压缩包条目数 = 5000
最大压缩比 = 1000


def _失败(错误码: str, 消息: str, *, 可重试: bool = False) -> 结果:
    return 结果.失败(错误码, 消息, 来源="演示文稿", 可重试=可重试)


def _文件摘要(路径: Path) -> str:
    return hashlib.sha256(路径.read_bytes()).hexdigest()


def _检查压缩包(路径: Path, 最大字节数: int) -> str | None:
    """OOXML 不可信 ZIP 检查：文件数/单项/总量/压缩比/路径逃逸/宏/外部文件引用。"""
    try:
        压缩包 = zipfile.ZipFile(路径)
    except (zipfile.BadZipFile, OSError, ValueError):
        return "文件损坏"
    try:
        条目表 = 压缩包.infolist()
        if len(条目表) > 最大压缩包条目数:
            return "超出限制"
        总量 = 0
        for 条目 in 条目表:
            名称 = 条目.filename.replace("\\", "/")
            if 条目.file_size > 最大字节数 or 条目.file_size > 最大压缩比 * max(条目.compress_size, 1):
                return "超出限制"
            总量 += 条目.file_size
            if 总量 > 最大字节数:
                return "超出限制"
            if 名称.startswith("/") or ".." in 名称.split("/") or 名称.lower() == "vbaproject.bin":
                return "文件损坏"
            if 名称.endswith(".rels"):
                try:
                    关系 = 压缩包.read(名称).decode("utf-8", "replace")
                except (KeyError, OSError, zipfile.BadZipFile, RuntimeError):
                    关系 = ""
                if 'TargetMode="External"' in 关系 and "file:" in 关系:
                    return "文件损坏"
        return None
    finally:
        压缩包.close()


def 解析演示文稿(
    文件路径: str | Path,
    格式: str = "",
    最大幻灯片数: int = 默认最大幻灯片数,
    最大字节数: int = 默认最大字节数,
    超时秒: float = 默认超时秒,
) -> 结果:
    """解析 ppt/pptx 演示文稿为 结果[通用文档]。"""
    开始时间 = time.monotonic()
    来源 = Path(文件路径)
    if not 来源.is_file():
        return _失败("文件不存在", f"文件不存在: {来源}")
    try:
        格式 = 归一化格式(格式 or 来源.suffix.lstrip("."))
    except ValueError:
        return _失败("参数不合法", f"不支持的格式: {格式 or 来源.suffix}")
    if 格式 not in ("ppt", "pptx"):
        return _失败("参数不合法", f"仅支持 ppt/pptx，收到: {格式}")
    if 来源.stat().st_size > 最大字节数:
        return _失败("超出限制", f"文件大小超过上限 {最大字节数} 字节")
    if pptx is None:
        return _失败("提供者不可用", "缺少 python-pptx，无法解析演示文稿", 可重试=True)
    if 格式 == "pptx":
        错误码 = _检查压缩包(来源, 最大字节数)
        if 错误码:
            return _失败(错误码, "OOXML 不可信 ZIP 检查未通过" if 错误码 == "文件损坏" else "演示文稿资源超过上限")
    临时目录 = None
    try:
        if 格式 == "ppt":
            临时目录 = Path(tempfile.mkdtemp(prefix="平台演示文稿_"))
            转换结果 = 转换办公文件(来源, "pptx", 输出目录=临时目录, 超时秒=超时秒, 最大输出字节=最大字节数)
            if not 转换结果.成功:
                return _失败(转换结果.错误码, f"ppt 经文档转换失败: {转换结果.错误说明}", 可重试=转换结果.可重试)
            解析路径, 解析方式 = Path(转换结果.值["路径"]), "转换"
        else:
            解析路径, 解析方式 = 来源, "原生"
        return _解析pptx(解析路径, 来源, 格式, 解析方式, 最大幻灯片数, 开始时间)
    finally:
        临时目录 and shutil.rmtree(临时目录, ignore_errors=True)


def _解析pptx(解析路径: Path, 原始路径: Path, 格式: str, 解析方式: str, 最大幻灯片数: int, 开始时间: float) -> 结果:
    """用 python-pptx 解析并组装通用文档。"""
    try:
        演示 = pptx.Presentation(str(解析路径))
    except Exception:
        return _失败("文件损坏", f"无法打开演示文稿（损坏或伪装）: {解析路径.name}")
    幻灯片数 = len(演示.slides)
    if 幻灯片数 > 最大幻灯片数:
        return _失败("超出限制", f"幻灯片数 {幻灯片数} 超过上限 {最大幻灯片数}")
    块列表, 资源列表, 警告 = [], [], []
    for 序号, 幻灯片 in enumerate(演示.slides, start=1):
        文本行, 附加, 图像引用 = [], {}, []
        for 形状 in 幻灯片.shapes:
            if 形状.has_text_frame:
                文本行 += [段落.text for 段落 in 形状.text_frame.paragraphs if 段落.text.strip()]
            if getattr(形状, "has_table", False):
                文本行 += [单元格.text for 行 in 形状.table.rows for 单元格 in 行.cells if 单元格.text.strip()]
            if 形状.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    图像 = 形状.image
                    资源索引 = len(资源列表)
                    资源列表.append(文档资源(类型="图像", 媒体类型=图像.content_type or "image/未知",
                        文件名=图像.filename or f"图像_{序号}_{资源索引}", 描述=f"幻灯片{序号} 图像",
                        字节数据b64=base64.b64encode(图像.blob).decode("ascii")))
                    图像引用.append(资源索引)
                except Exception as 错误:
                    警告.append(f"幻灯片{序号} 图像读取失败: {错误}")
        try:
            备注 = (幻灯片.notes_slide.notes_text_frame.text if 幻灯片.has_notes_slide else "").strip()
            备注 and 附加.update(备注=备注)
        except Exception as 错误:
            警告.append(f"幻灯片{序号} 备注读取失败: {错误}")
        if 图像引用:
            附加["图像资源引用"] = 图像引用
        块列表.append(文档块(类型="幻灯片", 文本="\n".join(文本行), 来源位置=来源位置(幻灯片=序号), 附加=附加))
    try:
        标题 = 演示.core_properties.title or 原始路径.stem
    except Exception:
        标题 = 原始路径.stem
    return 结果.成功结果(通用文档(
        文档类型="演示文稿", 格式=格式, 标题=标题, 块列表=块列表, 资源列表=资源列表,
        保真级别="高", 解析方式=解析方式, 警告=警告,
        诊断=[f"幻灯片数: {幻灯片数}", f"资源数: {len(资源列表)}"],
        耗时秒=round(time.monotonic() - 开始时间, 4),
        提供者版本={"python-pptx": getattr(pptx, "__version__", "未知")}, 原始文件摘要=_文件摘要(原始路径),
    ))
