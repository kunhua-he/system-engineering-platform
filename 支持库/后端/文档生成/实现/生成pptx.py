"""PPTX 原子生成：按幻灯片列表用 python-pptx 生成字节。

缺库抛 ImportError（上层转 提供者不可用）、参数非法抛 ValueError
（上层转 参数不合法）。
"""

from __future__ import annotations

import io


def 生成PPTX字节(参数: dict) -> bytes:
    """按 幻灯片列表 生成 PPTX 字节。"""
    from pptx import Presentation  # 缺库时 ImportError 冒泡
    from pptx.util import Inches

    幻灯片列表 = 参数.get("幻灯片列表") or []
    if not isinstance(幻灯片列表, list) or not 幻灯片列表:
        raise ValueError("幻灯片列表不能为空，至少需要一张幻灯片")

    演示文稿 = Presentation()
    演示文稿.slide_width = Inches(13.333)
    演示文稿.slide_height = Inches(7.5)
    渲染数 = 0
    for 幻灯片 in 幻灯片列表:
        if not isinstance(幻灯片, dict):
            continue
        页面 = 演示文稿.slides.add_slide(演示文稿.slide_layouts[1])
        标题占位 = 页面.shapes.title
        if 标题占位 is not None:
            标题占位.text = _取标题(幻灯片)
        要点 = _取要点(幻灯片)
        if 要点:
            _填充要点(页面, 要点)
        备注 = 幻灯片.get("备注") or 幻灯片.get("notes") or ""
        if 备注:
            页面.notes_slide.notes_text_frame.text = str(备注)
        渲染数 += 1

    if 渲染数 == 0:
        raise ValueError("幻灯片列表中没有可渲染幻灯片")

    缓冲 = io.BytesIO()
    演示文稿.save(缓冲)
    return 缓冲.getvalue()


def _取标题(幻灯片: dict) -> str:
    return str(
        幻灯片.get("标题")
        or 幻灯片.get("name")
        or 幻灯片.get("文本")
        or ""
    )


def _取要点(幻灯片: dict) -> list:
    要点 = 幻灯片.get("要点") or 幻灯片.get("bullets") or []
    if not isinstance(要点, list):
        return []
    结果: list[str] = []
    for 点 in 要点:
        if isinstance(点, dict):
            结果.append(str(点.get("文本", "")))
        else:
            结果.append(str(点))
    return 结果


def _填充要点(页面, 要点: list) -> None:
    """把要点文本写入正文占位符。"""
    正文 = 页面.placeholders[1]
    文本框 = 正文.text_frame
    文本框.word_wrap = True
    for 序号, 文本 in enumerate(要点):
        段落 = 文本框.paragraphs[0] if 序号 == 0 else 文本框.add_paragraph()
        段落.text = 文本
