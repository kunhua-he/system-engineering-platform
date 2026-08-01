"""python-docx 独立提供者：DOCX 解析与生成的唯一 python-docx 代码收拢点。

一第三方一支持库：本提供者只 import python-docx（纯 Python，主进程加载）。
解析：段落/表格/图像 → 通用文档字典；生成：内容参数 → docx 字节。
OOXML 按不可信 ZIP 校验。旧格式（doc）转换走 文档转换 模块（本提供者不处理）。
"""

from __future__ import annotations

import base64
import io
import zipfile
from pathlib import Path
from typing import Any

from 公共契约.基础类型.结果类型 import 结果

来源 = "python_docx提供者"
最大成员数 = 500
最大单项字节 = 50 * 1024 * 1024
最大解压总字节 = 200 * 1024 * 1024
最大压缩比 = 100.0


def _失败(错误码: str, 消息: str, *, 可重试: bool = False) -> 结果:
    return 结果.失败(错误码, 消息, 来源=来源, 可重试=可重试)


def _成功(值: Any) -> 结果:
    return 结果.成功结果(值)


def 校验OOXML安全(文件路径: Path) -> 结果:
    """不可信 ZIP 检查：文件数/单项大小/总大小/压缩比/路径逃逸/外部关系。"""
    try:
        with zipfile.ZipFile(str(文件路径)) as 压缩包:
            成员列表 = 压缩包.infolist()
            if not 成员列表:
                return _失败("文件损坏", "DOCX 压缩包内无任何成员")
            if len(成员列表) > 最大成员数:
                return _失败("超出限制", f"OOXML 成员数 {len(成员列表)} 超过上限 {最大成员数}")
            总解压字节 = 0
            原始总字节 = 0
            for 成员 in 成员列表:
                if 成员.file_size > 最大单项字节:
                    return _失败("超出限制", f"OOXML 单项 {成员.filename} 解压后超过上限")
                总解压字节 += 成员.file_size
                原始总字节 += 成员.compress_size
                if 总解压字节 > 最大解压总字节:
                    return _失败("超出限制", f"OOXML 总解压体积超过上限 {最大解压总字节}")
                规范化路径 = 成员.filename.replace("\\", "/")
                if ".." in 规范化路径.split("/") or 规范化路径.startswith("/"):
                    return _失败("文件损坏", f"OOXML 成员路径越界: {成员.filename}")
                if 成员.filename.lower().endswith(".rels"):
                    内容 = 压缩包.read(成员).decode("utf-8", errors="ignore")
                    if 'TargetMode="External"' in 内容:
                        return _失败("文件损坏", f"OOXML 含外部关系: {成员.filename}")
            if 原始总字节 > 0 and 总解压字节 / 原始总字节 > 最大压缩比:
                return _失败("超出限制", f"OOXML 压缩比 {总解压字节 / 原始总字节:.1f} 超过上限 {最大压缩比}")
    except zipfile.BadZipFile as 错误:
        return _失败("文件损坏", f"不是有效 OOXML 压缩包: {错误}")
    return _成功(None)


def 解析文字文档(文件路径: str, 格式: str = "docx", *,
              最大字节数: int = 200 * 1024 * 1024, 超时秒: float = 60) -> 结果:
    """解析 DOCX 为通用文档字典（段落/表格/图像资源）。"""
    if not isinstance(文件路径, str) or not 文件路径.strip():
        return _失败("参数不合法", "文件路径必须是非空文本")
    输入文件 = Path(文件路径)
    if not 输入文件.is_file():
        return _失败("文件不存在", f"文件不存在: {文件路径}")
    if 输入文件.stat().st_size > 最大字节数:
        return _失败("超出限制", f"文件过大: {输入文件.stat().st_size} 字节 > {最大字节数}")
    安全 = 校验OOXML安全(输入文件)
    if not 安全.成功:
        return 安全
    try:
        from docx import Document
        文档 = Document(str(输入文件))
    except Exception as 错误:
        return _失败("文件损坏", f"python-docx 打开失败: {错误}")
    块列表: list[dict] = []
    资源列表: list[dict] = []
    来源引用 = {"module": "docx-parser", "file_id": 0}
    # 段落
    for 段落 in 文档.paragraphs:
        if 段落.text.strip():
            块列表.append({"类型": "段落", "文本": 段落.text, "页码": 1,
                          "resource_ref": None, "source_ref": 来源引用})
    # 表格
    for 表格 in 文档.tables:
        表格数据 = [[单元格.text for 单元格 in 行.cells] for 行 in 表格.rows]
        块列表.append({"类型": "表格", "文本": "", "页码": 1,
                      "表格数据": 表格数据, "resource_ref": None,
                      "source_ref": 来源引用})
    # 图像
    try:
        for 关系 in 文档.part.rels.values():
            if "image" in str(关系.reltype):
                try:
                    图像字节 = 关系.target_part.blob
                    资源列表.append({
                        "类型": "image", "资源引用": 关系.rId,
                        "字节数据b64": base64.b64encode(图像字节).decode("ascii"),
                    })
                except Exception:
                    pass
    except Exception:
        pass
    通用文档 = {
        "schema_version": "通用文档-v1",
        "content_type": "word",
        "标题": 输入文件.stem,
        "块列表": 块列表,
        "资源列表": 资源列表,
        "元数据": {"来源": "python_docx提供者", "格式": "docx"},
        "警告": [],
        "resource_diagnostics": [],
        "附加": {},
    }
    return _成功(通用文档)


def _docx可用() -> bool:
    """python-docx 可用性检查（环境变量可注入禁用）。"""
    import os
    if os.environ.get("python_docx提供者_禁用库") == "docx":
        return False
    try:
        import docx  # noqa: F401
        return True
    except Exception:
        return False


def 生成文字文档(内容参数: dict) -> 结果:
    """按 内容参数（内容块列表：标题/段落）生成 DOCX 字节。"""
    if not isinstance(内容参数, dict):
        return _失败("参数不合法", "内容参数必须是字典")
    if not _docx可用():
        return _失败("提供者不可用", "python-docx 不可用")
    块列表 = 内容参数.get("内容块列表") or 内容参数.get("块列表") or []
    if not isinstance(块列表, list) or not 块列表:
        return _失败("参数不合法", "内容块列表 必须是非空列表")
    try:
        from docx import Document
        文档 = Document()
        for 块 in 块列表:
            if not isinstance(块, dict):
                continue
            类型 = 块.get("类型", "段落")
            文本 = str(块.get("文本", ""))
            if 类型 in ("标题", "heading"):
                文档.add_heading(文本, level=1)
            elif 类型 in ("段落", "paragraph"):
                文档.add_paragraph(文本)
        io流 = io.BytesIO()
        文档.save(io流)
        字节 = io流.getvalue()
        if not 字节:
            return _失败("生成失败", "python-docx 产出为空")
        import hashlib
        return _成功({
            "字节b64": base64.b64encode(字节).decode("ascii"),
            "媒体类型": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "摘要": hashlib.sha256(字节).hexdigest(),
            "字节数": len(字节), "格式": "docx", "诊断": [],
        })
    except Exception as 错误:
        return _失败("生成失败", f"python-docx 生成失败: {错误}")
