"""python-pptx 独立提供者：演示文稿解析与生成（仅 pptx，直接 import pptx，OOXML 按不可信 ZIP 处理）。"""

from __future__ import annotations
import base64, hashlib, io, time, zipfile
from pathlib import Path

from 公共契约.基础类型.文档结构 import 文档块, 文档资源, 来源位置, 通用文档, 生成产物
from 公共契约.基础类型.结果类型 import 结果

try:
    import pptx
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ImportError:
    pptx = MSO_SHAPE_TYPE = None

默认最大幻灯片数 = 200
默认最大字节数 = 200 * 1024 * 1024
最大压缩包条目数 = 5000
最大压缩比 = 1000
媒体类型pptx = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _失败(错误码: str, 消息: str, *, 可重试: bool = False) -> 结果:
    return 结果.失败(错误码, 消息, 来源="python_pptx提供者", 可重试=可重试)

def _检查压缩包(路径: Path, 最大字节数: int) -> str | None:
    try:
        压缩包 = zipfile.ZipFile(路径)
    except (zipfile.BadZipFile, OSError, ValueError):
        return "文件损坏"
    try:
        条目表 = 压缩包.infolist()
        if len(条目表) > 最大压缩包条目数:
            return "超出限制"
        总量 = 0
        for 条目 in 条目表:
            名称 = 条目.filename.replace("\\", "/")
            if 条目.file_size > 最大字节数 or 条目.file_size > 最大压缩比 * max(条目.compress_size, 1):
                return "超出限制"
            总量 += 条目.file_size
            if 总量 > 最大字节数 or 名称.startswith("/") or ".." in 名称.split("/") or 名称.lower() == "vbaproject.bin":
                return "超出限制" if 总量 > 最大字节数 else "文件损坏"
            if 名称.endswith(".rels"):
                关系 = ""
                try:
                    关系 = 压缩包.read(名称).decode("utf-8", "replace")
                except (KeyError, OSError, zipfile.BadZipFile, RuntimeError):
                    pass
                if 'TargetMode="External"' in 关系 and "file:" in 关系:
                    return "文件损坏"
        return None
    finally:
        压缩包.close()

def 解析演示文稿(
    文件路径: str | Path,
    格式: str = "",
    最大幻灯片数: int = 默认最大幻灯片数,
    最大字节数: int = 默认最大字节数,
    超时秒: float = 60,
) -> 结果:
    """解析 pptx 为 结果[通用文档]（幻灯片块/备注/图像资源）。超时秒保留供转换链兼容，原生解析不使用。"""
    开始时间 = time.monotonic()
    来源 = Path(文件路径)
    if not 来源.is_file():
        return _失败("文件不存在", f"文件不存在: {来源}")
    格式值 = (格式 or 来源.suffix.lstrip(".")).lower()
    if 格式值 != "pptx":
        return _失败("参数不合法", f"python-pptx 提供者仅支持 pptx，收到: {格式值 or 来源.suffix}")
    if 来源.stat().st_size > 最大字节数:
        return _失败("超出限制", f"文件大小超过上限 {最大字节数} 字节")
    if pptx is None:
        return _失败("提供者不可用", "缺少 python-pptx，无法解析演示文稿", 可重试=True)
    if (错误码 := _检查压缩包(来源, 最大字节数)):
        return _失败(错误码, "OOXML 不可信 ZIP 检查未通过" if 错误码 == "文件损坏" else "演示文稿资源超过上限")
    try:
        演示 = pptx.Presentation(str(来源))
    except Exception:
        return _失败("文件损坏", f"无法打开演示文稿（损坏或伪装）: {来源.name}")
    幻灯片数 = len(演示.slides)
    if 幻灯片数 > 最大幻灯片数:
        return _失败("超出限制", f"幻灯片数 {幻灯片数} 超过上限 {最大幻灯片数}")
    块列表, 资源列表, 警告 = [], [], []
    for 序号, 幻灯片 in enumerate(演示.slides, start=1):
        文本行, 附加, 图像引用 = [], {}, []
        for 形状 in 幻灯片.shapes:
            if 形状.has_text_frame:
                文本行 += [段落.text for 段落 in 形状.text_frame.paragraphs if 段落.text.strip()]
            if getattr(形状, "has_table", False):
                文本行 += [单元格.text for 行 in 形状.table.rows for 单元格 in 行.cells if 单元格.text.strip()]
            if 形状.shape_type == MSO_SHAPE_TYPE.PICTURE:
                try:
                    图像 = 形状.image
                    资源列表.append(文档资源(类型="图像", 媒体类型=图像.content_type or "image/未知", 文件名=图像.filename or f"图像_{序号}_{len(资源列表)}", 描述=f"幻灯片{序号} 图像", 字节数据b64=base64.b64encode(图像.blob).decode("ascii")))
                    图像引用.append(len(资源列表) - 1)
                except Exception as 错误:
                    警告.append(f"幻灯片{序号} 图像读取失败: {错误}")
        try:
            备注 = (幻灯片.notes_slide.notes_text_frame.text if 幻灯片.has_notes_slide else "").strip()
            备注 and 附加.update(备注=备注)
        except Exception as 错误:
            警告.append(f"幻灯片{序号} 备注读取失败: {错误}")
        if 图像引用:
            附加["图像资源引用"] = 图像引用
        块列表.append(文档块(类型="幻灯片", 文本="\n".join(文本行), 来源位置=来源位置(幻灯片=序号), 附加=附加))
    try:
        标题 = 演示.core_properties.title or 来源.stem
    except Exception:
        标题 = 来源.stem
    return 结果.成功结果(通用文档(文档类型="演示文稿", 格式="pptx", 标题=标题, 块列表=块列表, 资源列表=资源列表, 保真级别="高", 解析方式="原生", 警告=警告, 诊断=[f"幻灯片数: {幻灯片数}", f"资源数: {len(资源列表)}"], 耗时秒=round(time.monotonic() - 开始时间, 4), 提供者版本={"python-pptx": getattr(pptx, "__version__", "未知")}, 原始文件摘要=hashlib.sha256(来源.read_bytes()).hexdigest()))

def 生成演示文稿(内容参数: dict) -> 结果:
    """按 幻灯片列表 生成 pptx，返回 结果[生成产物字典]（字节b64/媒体类型/摘要）。"""
    开始时间 = time.monotonic()
    if not isinstance(内容参数, dict):
        return _失败("参数不合法", "内容参数必须是字典")
    if pptx is None:
        return _失败("提供者不可用", "缺少 python-pptx，无法生成演示文稿", 可重试=True)
    try:
        字节 = _生成字节(内容参数)
    except ValueError as 错误:
        return _失败("参数不合法", str(错误))
    except Exception as 错误:
        return _失败("生成失败", f"pptx 生成异常: {错误}")
    if not 字节:
        return _失败("生成失败", "pptx 生成结果为空")
    耗时秒 = round(time.monotonic() - 开始时间, 4)
    产物 = 生成产物(格式="pptx", 字节=字节, 媒体类型=媒体类型pptx, 摘要=hashlib.sha256(字节).hexdigest(), 诊断=[f"pptx 生成成功，耗时 {耗时秒} 秒"], 附加={"提供者版本": {"python-pptx": getattr(pptx, "__version__", "未知")}, "生成耗时秒": 耗时秒})
    return 结果.成功结果({**产物.转字典(), "字节b64": base64.b64encode(字节).decode("ascii")})

def _生成字节(内容参数: dict) -> bytes:
    from pptx.util import Inches
    幻灯片列表 = 内容参数.get("幻灯片列表") or []
    if not isinstance(幻灯片列表, list) or not 幻灯片列表:
        raise ValueError("幻灯片列表不能为空，至少需要一张幻灯片")
    演示文稿 = pptx.Presentation()
    演示文稿.slide_width = Inches(13.333)
    演示文稿.slide_height = Inches(7.5)
    渲染数 = 0
    for 幻灯片 in 幻灯片列表:
        if not isinstance(幻灯片, dict):
            continue
        页面 = 演示文稿.slides.add_slide(演示文稿.slide_layouts[1])
        if (标题占位 := 页面.shapes.title) is not None:
            标题占位.text = str(幻灯片.get("标题") or 幻灯片.get("name") or 幻灯片.get("文本") or "")
        if isinstance(要点 := 幻灯片.get("要点") or 幻灯片.get("bullets") or [], list) and 要点:
            正文 = 页面.placeholders[1].text_frame
            正文.word_wrap = True
            for 序号, 文本 in enumerate(要点):
                段落 = 正文.paragraphs[0] if 序号 == 0 else 正文.add_paragraph()
                段落.text = str(文本.get("文本", "")) if isinstance(文本, dict) else str(文本)
        if 备注 := 幻灯片.get("备注") or 幻灯片.get("notes") or "":
            页面.notes_slide.notes_text_frame.text = str(备注)
        渲染数 += 1
    if 渲染数 == 0:
        raise ValueError("幻灯片列表中没有可渲染幻灯片")
    缓冲 = io.BytesIO()
    演示文稿.save(缓冲)
    return 缓冲.getvalue()
