"""openpyxl 提供者真实测试：生成→解析往返/多工作表中文/公式值策略/损坏/伪装/加密/空表/缺提供者/超限/注册能力。"""

from __future__ import annotations

import base64, hashlib, io, sys, tempfile, unittest, zipfile
from pathlib import Path
from unittest import mock

if str(Path(__file__).resolve().parents[2]) not in sys.path:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from 支持库.适配层.openpyxl提供者 import 生成表格文档, 解析表格文档, 注册能力
from 支持库.适配层.openpyxl提供者.实现 import 表格文档 as 解析模块


def _造公式文件(路径: Path) -> None:
    from openpyxl import Workbook
    工作簿 = Workbook()
    工作表 = 工作簿.active
    工作表["A1"] = 10
    工作表["A2"] = 20
    工作表["A3"] = "=SUM(A1:A2)"
    工作簿.save(str(路径))


class TestOpenpyxl提供者(unittest.TestCase):
    def setUp(self):
        self.临时目录 = Path(tempfile.mkdtemp(prefix="测试_openpyxl提供者_"))

    def tearDown(self):
        import shutil
        shutil.rmtree(self.临时目录, ignore_errors=True)

    def test_生成与解析往返_多工作表中文(self):
        结果 = 生成表格文档({"工作表列表": [
            {"表名": "成绩单", "列": ["姓名", "分数", "备注"], "行": [["张三", 95, "优秀"], ["李四", 88, "良好"]]},
            {"表名": "汇总", "行": [["合计"], [183]]},
        ]})
        self.assertTrue(结果.成功, 结果.错误说明)
        产物 = 结果.值
        self.assertEqual(产物["格式"], "xlsx")
        self.assertEqual(产物["媒体类型"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        字节 = base64.b64decode(产物["字节b64"])
        self.assertEqual(len(字节), 产物["字节数"])
        self.assertEqual(hashlib.sha256(字节).hexdigest(), 产物["摘要"])
        (self.临时目录 / "往返.xlsx").write_bytes(字节)
        解析结果 = 解析表格文档(str(self.临时目录 / "往返.xlsx"))
        self.assertTrue(解析结果.成功, 解析结果.错误说明)
        文档 = 解析结果.值
        self.assertEqual((文档["文档类型"], 文档["格式"]), ("表格", "xlsx"))
        self.assertEqual([块["来源位置"]["工作表"] for 块 in 文档["块列表"]], ["成绩单", "汇总"])
        首块 = 文档["块列表"][0]
        self.assertEqual(首块["表格数据"], [["姓名", "分数", "备注"], ["张三", "95", "优秀"], ["李四", "88", "良好"]])
        for 键 in ("文档类型", "格式", "标题", "块列表", "资源列表", "保真级别", "解析方式", "提供者版本", "原始文件摘要"):
            self.assertIn(键, 文档)
        self.assertIn("openpyxl", 文档["提供者版本"])
        self.assertEqual(文档["附加"]["工作表数"], 2)

    def test_生成产物可重新打开(self):
        结果 = 生成表格文档({"工作表列表": [{"表名": "数据", "行": [[{"文本": "甲"}, 1]]}]})
        self.assertTrue(结果.成功)
        from openpyxl import load_workbook
        工作簿 = load_workbook(io.BytesIO(base64.b64decode(结果.值["字节b64"])), read_only=True)
        try:
            self.assertEqual(工作簿.sheetnames, ["数据"])
            self.assertEqual([行 for 行 in 工作簿["数据"].iter_rows(values_only=True)], [("甲", "1")])
        finally:
            工作簿.close()

    def test_公式值策略(self):
        路径 = self.临时目录 / "公式.xlsx"
        _造公式文件(路径)
        真模式 = 解析表格文档(str(路径), 数据模式=True)
        self.assertTrue(真模式.成功)
        self.assertEqual(真模式.值["块列表"][0]["表格数据"], [["10"], ["20"]])
        假模式 = 解析表格文档(str(路径), 数据模式=False)
        self.assertTrue(假模式.成功)
        self.assertEqual(假模式.值["块列表"][0]["表格数据"], [["10"], ["20"], ["=SUM(A1:A2)"]])

    def test_空行过滤与空表(self):
        from openpyxl import Workbook
        空行路径 = self.临时目录 / "空行.xlsx"
        工作表 = Workbook().active
        工作表["A1"] = 1
        工作表["A3"] = 3
        工作表.parent.save(str(空行路径))
        结果 = 解析表格文档(str(空行路径))
        self.assertTrue(结果.成功)
        self.assertEqual(结果.值["块列表"][0]["表格数据"], [["1"], ["3"]])
        空表路径 = self.临时目录 / "空表.xlsx"
        Workbook().save(str(空表路径))
        空表结果 = 解析表格文档(str(空表路径))
        self.assertTrue(空表结果.成功)
        self.assertEqual(len(空表结果.值["块列表"]), 1)
        self.assertEqual(空表结果.值["块列表"][0]["表格数据"], [])

    def test_文件不存在与参数不合法(self):
        结果 = 解析表格文档("/不存在的路径/文件.xlsx")
        self.assertEqual(结果.错误码, "文件不存在")
        路径 = self.临时目录 / "任意.xlsx"
        路径.write_bytes(b"x")
        非法 = 解析表格文档(str(路径), 数据模式="yes")
        self.assertEqual(非法.错误码, "参数不合法")

    def test_损坏与伪装与缺workbook(self):
        损坏 = self.临时目录 / "损坏.xlsx"
        损坏.write_bytes(b"not a zip" * 50)
        self.assertEqual(解析表格文档(str(损坏)).错误码, "文件损坏")
        伪装 = self.临时目录 / "伪装.xlsx"
        伪装.write_text("这是文本", encoding="utf-8")
        self.assertEqual(解析表格文档(str(伪装)).错误码, "文件损坏")
        缺 = self.临时目录 / "缺workbook.xlsx"
        with zipfile.ZipFile(str(缺), "w") as 压缩包:
            压缩包.writestr("随便.txt", "内容")
        self.assertEqual(解析表格文档(str(缺)).错误码, "文件损坏")

    def test_加密文件(self):
        路径 = self.临时目录 / "加密.xlsx"
        _造公式文件(路径)
        from openpyxl.utils.exceptions import InvalidFileException
        with mock.patch.object(解析模块, "_解析xlsx内容", side_effect=InvalidFileException("File contains an encrypted package")):
            结果 = 解析表格文档(str(路径))
        self.assertEqual(结果.错误码, "文件加密")

    def test_缺提供者(self):
        不可用 = {"openpyxl": None, "版本": {"openpyxl": "不可用"}}
        路径 = self.临时目录 / "任意.xlsx"
        路径.write_bytes(b"x")
        with mock.patch.object(解析模块, "加载提供者", return_value=不可用):
            解析结果 = 解析表格文档(str(路径))
            生成结果 = 生成表格文档({"工作表列表": [{"表名": "甲", "行": [[1]]}]})
        self.assertEqual(解析结果.错误码, "提供者不可用")
        self.assertEqual(生成结果.错误码, "提供者不可用")

    def test_超出限制(self):
        路径 = self.临时目录 / "大文件.xlsx"
        _造公式文件(路径)
        结果 = 解析表格文档(str(路径), 最大字节数=10)
        self.assertEqual(结果.错误码, "超出限制")
        截断 = 解析表格文档(str(路径), 最大工作表数=0)
        self.assertTrue(截断.成功)
        self.assertEqual(截断.值["块列表"], [])

    def test_生成参数不合法(self):
        for 参数 in ("不是字典", {"工作表列表": []}, {"工作表列表": [{"表名": "空表"}]}):
            self.assertEqual(生成表格文档(参数).错误码, "参数不合法")

    def test_注册能力(self):
        from 公共契约.能力契约.契约 import 能力注册表
        注册表 = 能力注册表()
        注册能力(注册表)
        self.assertEqual(注册表.能力id列表, ["表格文档.生成表格文档", "表格文档.解析表格文档"])
        实现 = 注册表.获取("表格文档.解析表格文档")
        self.assertEqual([参数["名称"] for 参数 in 实现.参数], ["文件路径", "数据模式"])
        调用结果 = 实现.调用("/不存在的路径/文件.xlsx")
        self.assertEqual(调用结果.错误码, "文件不存在")


if __name__ == "__main__":
    unittest.main()
