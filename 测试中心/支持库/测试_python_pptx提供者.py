"""python-pptx 提供者真实测试：生成→解析往返/中文/多页/文本框/图像/损坏/空演示/缺提供者/注册能力。"""

from __future__ import annotations
import base64, importlib.util, io, shutil, sys, tempfile, unittest, zipfile
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))
提供者目录 = (Path(__file__).resolve().parents[2] / "支持库" / "后端"
              / "文档转换支持库" / "python_pptx提供者")


def 加载模块(名称: str, 相对路径: str):
    规格 = importlib.util.spec_from_file_location(名称, 提供者目录 / 相对路径)
    模块 = importlib.util.module_from_spec(规格)
    规格.loader.exec_module(模块)
    return 模块

实现模块 = 加载模块("python_pptx提供者_实现", "实现/演示文稿.py")
入口模块 = 加载模块("python_pptx提供者_入口", "__init__.py")
最小PNG = base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==")

def 生成夹具(路径: Path) -> None:
    from pptx import Presentation
    from pptx.util import Inches

    演示 = Presentation()
    封面 = 演示.slides.add_slide(演示.slide_layouts[0])
    封面.shapes.title.text = "中文演示文稿测试"
    封面.placeholders[1].text_frame.text = "第一要点：系统级支持库\n第二要点：python-pptx 中文适配"
    封面.notes_slide.notes_text_frame.text = "这是封面备注文本"
    for 序号 in (2, 3):
        页 = 演示.slides.add_slide(演示.slide_layouts[1])
        页.shapes.title.text = f"第{序号}页标题"
        框 = 页.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2)); 框.text_frame.text = f"第{序号}页文本框内容"
    图像页 = 演示.slides.add_slide(演示.slide_layouts[6])
    图像页.shapes.add_picture(io.BytesIO(最小PNG), Inches(1), Inches(1), Inches(2), Inches(2))
    演示.save(str(路径))


class Test解析演示文稿(unittest.TestCase):
    def setUp(self):
        self.临时目录 = Path(tempfile.mkdtemp(prefix="测试_python_pptx提供者_"))
        self.夹具路径 = self.临时目录 / "夹具.pptx"
        生成夹具(self.夹具路径)

    def tearDown(self):
        shutil.rmtree(self.临时目录, ignore_errors=True)

    def test_多页中文文本框往返(self):
        结果 = 实现模块.解析演示文稿(self.夹具路径)
        self.assertTrue(结果.成功, 结果.错误说明)
        文档 = 结果.值
        self.assertEqual((文档.文档类型, 文档.格式, 文档.解析方式), ("演示文稿", "pptx", "原生"))
        self.assertEqual(len(文档.块列表), 4)
        首块 = 文档.块列表[0]
        self.assertEqual((首块.类型, 首块.来源位置.幻灯片), ("幻灯片", 1))
        self.assertIn("中文演示文稿测试", 首块.文本)
        self.assertIn("python-pptx 中文适配", 首块.文本)
        self.assertEqual((首块.附加["备注"], 文档.块列表[3].附加["图像资源引用"]), ("这是封面备注文本", [0]))
        self.assertEqual(文档.块列表[1].来源位置.幻灯片, 2)
        self.assertIn("第2页标题", 文档.块列表[1].文本)
        self.assertIn("第2页文本框内容", 文档.块列表[1].文本)
        self.assertIn("第3页文本框内容", 文档.块列表[2].文本)
        图像 = 文档.资源列表[0]
        self.assertEqual((图像.类型, 图像.字节数据b64), ("图像", base64.b64encode(最小PNG).decode("ascii")))
        self.assertTrue(图像.媒体类型.startswith("image/"))
        self.assertEqual(len(文档.原始文件摘要), 64)

    def test_转字典结构兼容(self):
        字典 = 实现模块.解析演示文稿(self.夹具路径).值.转字典()
        首块 = 字典["块列表"][0]
        self.assertEqual((字典["文档类型"], 字典["格式"], 首块["类型"], 首块["来源位置"]["幻灯片"], "中文演示文稿测试" in 首块["文本"]), ("演示文稿", "pptx", "幻灯片", 1, True))

    def test_文件不存在与格式非法(self):
        self.assertEqual(实现模块.解析演示文稿(self.临时目录 / "不存在.pptx").错误码, "文件不存在")
        for 格式 in ("docx", "ppt", "pdf"):
            self.assertEqual(实现模块.解析演示文稿(self.夹具路径, 格式=格式).错误码, "参数不合法", 格式)

    def test_损坏与伪装(self):
        原始 = self.夹具路径.read_bytes()
        损坏路径 = self.临时目录 / "损坏.pptx"
        损坏路径.write_bytes(原始[: len(原始) // 2])
        self.assertEqual(实现模块.解析演示文稿(损坏路径).错误码, "文件损坏")
        伪装1 = self.临时目录 / "伪装1.pptx"
        伪装1.write_bytes("这不是压缩包内容".encode("utf-8") * 100)
        self.assertEqual(实现模块.解析演示文稿(伪装1).错误码, "文件损坏")
        伪装2 = self.临时目录 / "伪装2.pptx"
        with zipfile.ZipFile(伪装2, "w") as 包:
            包.writestr("随便.txt", "不是演示文稿")
        self.assertEqual(实现模块.解析演示文稿(伪装2).错误码, "文件损坏")

    def test_超幻灯片数与超大小(self):
        self.assertEqual(实现模块.解析演示文稿(self.夹具路径, 最大幻灯片数=2).错误码, "超出限制")
        self.assertEqual(实现模块.解析演示文稿(self.夹具路径, 最大字节数=100).错误码, "超出限制")

    def test_空演示解析成功零块(self):
        from pptx import Presentation
        Presentation().save(str(self.临时目录 / "空.pptx"))
        文档 = 实现模块.解析演示文稿(self.临时目录 / "空.pptx").值
        self.assertEqual((文档.块列表, 文档.资源列表), ([], []))

    def test_缺python_pptx提供者不可用(self):
        with mock.patch.object(实现模块, "pptx", None):
            结果 = 实现模块.解析演示文稿(self.夹具路径)
        self.assertEqual(结果.错误码, "提供者不可用")


class Test生成演示文稿(unittest.TestCase):
    def setUp(self):
        self.临时目录 = Path(tempfile.mkdtemp(prefix="测试_python_pptx提供者_生成_"))

    def tearDown(self):
        shutil.rmtree(self.临时目录, ignore_errors=True)

    def test_生成并解析往返中文多页(self):
        结果 = 实现模块.生成演示文稿({"幻灯片列表": [
            {"标题": "中文演示文稿测试", "要点": ["第一要点：系统级支持库", "第二要点：python-pptx 中文适配"], "备注": "这是封面备注"},
            {"标题": "第二页", "要点": ["正文要点A"], "备注": ""},
        ]})
        self.assertTrue(结果.成功, 结果.错误说明)
        产物 = 结果.值
        self.assertEqual((产物["格式"], 产物["媒体类型"], len(产物["摘要"]), "字节b64" in 产物), ("pptx", 实现模块.媒体类型pptx, 64, True))
        字节 = base64.b64decode(产物["字节b64"])
        self.assertGreater(len(字节), 1000)
        路径 = self.临时目录 / "往返.pptx"
        路径.write_bytes(字节)
        文档 = 实现模块.解析演示文稿(路径).值
        self.assertEqual(len(文档.块列表), 2)
        self.assertEqual(文档.块列表[0].来源位置.幻灯片, 1)
        self.assertIn("中文演示文稿测试", 文档.块列表[0].文本)
        self.assertIn("python-pptx 中文适配", 文档.块列表[0].文本)
        self.assertEqual(文档.块列表[0].附加["备注"], "这是封面备注")
        self.assertEqual(文档.块列表[1].来源位置.幻灯片, 2)
        self.assertIn("第二页", 文档.块列表[1].文本)

    def test_内容参数不合法(self):
        for 参数 in ("不是字典", {}, {"幻灯片列表": []}, {"幻灯片列表": ["不是字典", 42]}):
            self.assertEqual(实现模块.生成演示文稿(参数).错误码, "参数不合法")

    def test_缺python_pptx生成提供者不可用(self):
        with mock.patch.object(实现模块, "pptx", None):
            结果 = 实现模块.生成演示文稿({"幻灯片列表": [{"标题": "示例标题"}]})
        self.assertEqual(结果.错误码, "提供者不可用")


class Test注册能力(unittest.TestCase):
    def test_入口注册内部能力且阻断公开owner(self):
        记录 = []

        class 假注册表:
            def 注册(self, 能力):
                记录.append(能力)

        入口模块.注册能力(假注册表())
        能力id列表 = [能力.能力id for 能力 in 记录]
        self.assertEqual(能力id列表, ["内部.演示文稿.解析", "内部.演示文稿.生成"])
        # Provider 不能夺取公开 owner；公开能力由 支持库/后端/演示文稿
        # 与 支持库/后端/文档生成 分别持有。
        self.assertNotIn("办公文档支持库.演示文稿.解析演示文稿", 能力id列表)
        self.assertNotIn("办公文档支持库.演示文稿.生成演示文稿", 能力id列表)
        self.assertEqual(([参数["名称"] for 参数 in 记录[0].参数], [参数["名称"] for 参数 in 记录[1].参数]), (["文件路径", "格式", "最大幻灯片数", "最大字节数", "超时秒"], ["内容参数"]))


if __name__ == "__main__":
    unittest.main()
