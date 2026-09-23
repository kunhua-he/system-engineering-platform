"""演示文稿支持库真实测试：pptx 生成→解析往返、损坏、伪装、缺提供者、ppt 转换链。"""

from __future__ import annotations

import base64
import contextlib
import importlib
import io
import os
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
import zipfile
from pathlib import Path
from unittest import mock

if str(Path(__file__).resolve().parents[2]) not in sys.path:
    sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

系统根 = Path(__file__).resolve().parents[2]

from 公共契约.基础类型.结果类型 import 结果
from 公共契约.运行时 import 进程终止, 平台适配
from 支持库.后端.办公文档支持库.演示文稿 import 解析演示文稿

#: ★ 受管临时根（仓库内**固定排除目录** `工程缓存/` 下）—— LibreOffice **私有档案**的落点。
#: 2026-09-23 收口：原写法把 `dir=` 指向 `输出目录.parent`，而 `输出目录` 是**函数形参**，
#: 判据一静态解析不出（落进 `测试写入-写动作未解析`）；且裸 `mkdtemp()` 的落点由测试运行时的
#: `TMPDIR` 决定（平台跑测试时被指进仓库工作目录）。显式指向 `工程缓存/` 后：落点与运行环境
#: 解耦、判据一判「受管写」，即便被 SIGKILL 残留也进不了工作区指纹。档案由 `finally` 里的
#: `shutil.rmtree` 清理。
受管临时根 = 系统根 / "工程缓存" / "测试临时"
受管临时根.mkdir(parents=True, exist_ok=True)
from 支持库.后端.文档转换支持库.LibreOffice转换 import 检查提供者

默认soffice路径 = "/Applications/LibreOffice.app/Contents/MacOS/soffice"


def _找soffice() -> str | None:
    """定位 soffice（环境变量 → PATH → macOS 应用包内路径），找不到返回 None。"""
    候选 = [os.getenv("LIBREOFFICE_BIN"), os.getenv("SOFFICE_BIN"),
           shutil.which("soffice"), 默认soffice路径]
    return next((项 for 项 in 候选 if 项 and Path(项).exists()), None)


def _转换办公文件(源: Path, 输出目录: Path, 目标格式: str, *,
              超时秒: float = 120.0, 尝试次数: int = 3) -> Path | None:
    """经 soffice 把 源 转成 目标格式，返回产物路径；重试后仍无产出返回 None。

    两处并发脆弱点在此收口：
    1) 私有档案：LibreOffice 桌面端是单实例，同一 UserInstallation（默认落在用户档案
       目录）下后到的 --convert-to 会被静默丢弃——退出码 0 但零产出。故每次尝试都用
       独立私有档案，与同机其它进程的实例完全隔离。
    2) 有限次重试：仍按指数退避重试，避免单次竞态直接把用例判红。
    """
    soffice = _找soffice()
    if not soffice:
        return None
    目标 = 输出目录 / f"{源.stem}.{目标格式}"
    for 第几次 in range(1, 尝试次数 + 1):
        档案 = Path(tempfile.mkdtemp(prefix=f"私有档案_{第几次}_", dir=受管临时根))
        try:
            # 目标与源不同名、且每轮先清目标：LibreOffice 遇到已存在的同名目标会拒绝覆盖
            if 目标.exists():
                目标.unlink()
            进程 = subprocess.Popen(
                [soffice, f"-env:UserInstallation=file://{档案}", "--headless",
                 "--norestore", "--convert-to", 目标格式,
                 "--outdir", str(输出目录), str(源)],
                stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                **平台适配.子进程组启动标志())
            try:
                进程.communicate(timeout=超时秒)
            except subprocess.TimeoutExpired:
                # 按进程组回收：子 shell 包一层时也不会留下孤儿 soffice 持有档案锁
                进程终止.强制结束子进程(进程, 宽限秒=1.0, 等待秒=2.0)
                进程.communicate()
            if 目标.is_file() and 目标.stat().st_size > 0:
                return 目标
        finally:
            shutil.rmtree(档案, ignore_errors=True)
        if 第几次 < 尝试次数:
            time.sleep(2 ** (第几次 - 1))
    return None


@contextlib.contextmanager
def _实现走私有档案(档案根: Path):
    """让被测链路（实现内部的 soffice 调用）也走私有档案。

    实现 支持库.后端.文档转换支持库.LibreOffice转换 只用 查找LibreOffice() 取可执行
    路径（并缓存）。这里把该函数指向一个「每次调用用独立私有档案」的包装脚本——
    只换可执行入口，转换本身仍是真实 soffice，不做任何打桩假成功。
    """
    from 支持库.后端.文档转换支持库.LibreOffice转换.实现 import 文档转换 as 文档转换模块

    真实 = _找soffice()
    if not 真实:
        yield None
        return
    档案根.mkdir(parents=True, exist_ok=True)
    包装 = 档案根 / "soffice私有档案"
    # 脚本内变量名必须用 ASCII：macOS /bin/sh 不认多字节变量名（会报 Illegal byte sequence）
    包装.write_text(
        "#!/bin/sh\n"
        f'prof="{档案根}/call_$$"\n'
        'mkdir -p "$prof" || exit 1\n'
        "trap 'rm -rf \"$prof\"' EXIT\n"
        f'"{真实}" -env:UserInstallation="file://$prof" "$@"\n'
        'rc=$?\n'
        'exit $rc\n')
    os.chmod(包装, 0o755)
    with mock.patch.object(文档转换模块, "查找LibreOffice", return_value=str(包装)):
        yield str(包装)


class 假调用器:
    """测试注入的假能力调用器：所有能力返回 提供者不可用。"""

    def 调用能力(self, 能力id, 参数=None, **关键字):
        return 结果.失败("提供者不可用", f"{能力id} 不可用（模拟调用器）", 来源="测试", 可重试=True)

    def 幂等重放(self, *args, **kwargs):
        return False

    def 查询调用历史(self, 上限=50):
        return []

    def 最近失败(self, 上限=10):
        return []

    def 回答九问(self, *args, **kwargs):
        return {}

最小PNG = base64.b64decode(
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg=="
)


def 生成示例演示文稿(路径: Path) -> None:
    """python-pptx 生成多页演示文稿：封面(标题+要点+中文+备注) + 正文页 + 图像页。"""
    from pptx import Presentation
    from pptx.util import Inches

    演示 = Presentation()
    封面 = 演示.slides.add_slide(演示.slide_layouts[0])
    封面.shapes.title.text = "中文演示文稿测试"
    正文 = 封面.placeholders[1].text_frame
    正文.text = "第一要点：系统级支持库"
    段落 = 正文.add_paragraph()
    段落.text = "第二要点：python-pptx 中文适配"
    封面.notes_slide.notes_text_frame.text = "这是封面备注文本"
    for 序号 in (2, 3):
        页 = 演示.slides.add_slide(演示.slide_layouts[1])
        页.shapes.title.text = f"第{序号}页标题"
        框 = 页.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(2))
        框.text_frame.text = f"第{序号}页正文内容"
    图像页 = 演示.slides.add_slide(演示.slide_layouts[6])
    图像页.shapes.add_picture(io.BytesIO(最小PNG), Inches(1), Inches(1), Inches(2), Inches(2))
    演示.save(str(路径))


class Test演示文稿pptx(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        """装配唯一能力调用服务：安装全部支持库（含受管提供者）并绑定。"""
        from 公共契约.能力契约.契约 import 能力注册表
        from 运行核心.能力调用.唯一能力调用 import 创建并绑定
        from 运行核心.加载器.包安装.支持库安装 import 安装全部支持库

        cls.注册表 = 能力注册表()
        安装全部支持库(系统根 / "支持库", cls.注册表)
        cls.服务 = 创建并绑定(cls.注册表)

    @classmethod
    def tearDownClass(cls):
        from 运行核心.能力调用.唯一能力调用 import 销毁全局唯一服务
        销毁全局唯一服务()

    def setUp(self):
        self.临时目录 = Path(tempfile.mkdtemp(prefix="测试_演示文稿_"))
        self.pptx路径 = self.临时目录 / "示例演示.pptx"
        生成示例演示文稿(self.pptx路径)

    def tearDown(self):
        shutil.rmtree(self.临时目录, ignore_errors=True)

    def test_生成解析往返(self):
        结果 = 解析演示文稿(self.pptx路径)
        self.assertTrue(结果.成功, 结果.错误说明)
        文档 = 结果.值
        self.assertEqual(文档.文档类型, "演示文稿")
        self.assertEqual(文档.格式, "pptx")
        self.assertEqual(文档.解析方式, "原生")
        self.assertEqual(文档.保真级别, "高")
        self.assertEqual(len(文档.块列表), 4)
        首块 = 文档.块列表[0]
        self.assertEqual(首块.类型, "幻灯片")
        self.assertEqual(首块.来源位置.幻灯片, 1)
        self.assertIn("中文演示文稿测试", 首块.文本)
        self.assertIn("第一要点", 首块.文本)
        self.assertIn("python-pptx 中文适配", 首块.文本)
        self.assertEqual(首块.附加["备注"], "这是封面备注文本")
        self.assertEqual(文档.块列表[1].来源位置.幻灯片, 2)
        self.assertIn("第2页正文内容", 文档.块列表[1].文本)
        self.assertEqual(len(文档.资源列表), 1)
        图像 = 文档.资源列表[0]
        self.assertEqual(图像.类型, "图像")
        self.assertTrue(图像.媒体类型.startswith("image/"))
        self.assertEqual(图像.字节数据b64, base64.b64encode(最小PNG).decode("ascii"))
        self.assertEqual(文档.块列表[3].附加["图像资源引用"], [0])
        self.assertTrue(文档.提供者版本.get("python-pptx"))
        self.assertEqual(len(文档.原始文件摘要), 64)
        self.assertGreaterEqual(文档.耗时秒, 0)

    def test_超幻灯片数(self):
        结果 = 解析演示文稿(self.pptx路径, 最大幻灯片数=2)
        self.assertFalse(结果.成功)
        self.assertEqual(结果.错误码, "超出限制")

    def test_超文件大小(self):
        结果 = 解析演示文稿(self.pptx路径, 最大字节数=100)
        self.assertFalse(结果.成功)
        self.assertEqual(结果.错误码, "超出限制")

    def test_格式参数不合法(self):
        for 格式 in ("docx", "pdf"):
            结果 = 解析演示文稿(self.pptx路径, 格式=格式)
            self.assertEqual(结果.错误码, "参数不合法", 格式)

    def test_文件不存在(self):
        结果 = 解析演示文稿(self.临时目录 / "不存在.pptx")
        self.assertEqual(结果.错误码, "文件不存在")

    def test_损坏文件(self):
        原始 = self.pptx路径.read_bytes()
        损坏路径 = self.临时目录 / "损坏.pptx"
        损坏路径.write_bytes(原始[: len(原始) // 2])
        结果 = 解析演示文稿(损坏路径)
        self.assertEqual(结果.错误码, "文件损坏")

    def test_扩展名伪装(self):
        伪装1 = self.临时目录 / "伪装1.pptx"
        伪装1.write_bytes("这不是压缩包内容".encode("utf-8") * 100)
        self.assertEqual(解析演示文稿(伪装1).错误码, "文件损坏")
        伪装2 = self.临时目录 / "伪装2.pptx"
        with zipfile.ZipFile(伪装2, "w") as 包:
            包.writestr("随便.txt", "不是演示文稿")
        self.assertEqual(解析演示文稿(伪装2).错误码, "文件损坏")

    def test_缺python_pptx返回提供者不可用(self):
        # 临时注入假调用器（受管提供者不可用）→ 如实返回 提供者不可用
        from 运行核心.能力调用.唯一能力调用 import 创建并绑定, 设置全局唯一服务

        设置全局唯一服务(假调用器())
        try:
            结果 = 解析演示文稿(self.pptx路径)
            self.assertEqual(结果.错误码, "提供者不可用")
        finally:
            创建并绑定(self.__class__.注册表)


class Test演示文稿ppt转换链(unittest.TestCase):
    def setUp(self):
        self.临时目录 = Path(tempfile.mkdtemp(prefix="测试_演示文稿_ppt_"))
        self.pptx路径 = self.临时目录 / "示例.pptx"
        生成示例演示文稿(self.pptx路径)

    def tearDown(self):
        shutil.rmtree(self.临时目录, ignore_errors=True)

    def test_pptx转ppt再经转换链解析(self):
        状态 = 检查提供者()
        self.assertTrue(状态.成功)
        self.assertEqual(状态.值.get("LibreOffice"), "可用", "本机需装 LibreOffice 才能真实转换")
        # 夹具：平台文档转换暂不支持生成 ppt，直接用 LibreOffice 把示例 pptx 转成旧版 .ppt
        # （私有档案 + 有限次重试：并发下默认档案单实例会让后到者「退出码 0 零产出」）
        输出目录 = self.临时目录 / "转ppt"
        输出目录.mkdir()
        ppt路径 = _转换办公文件(self.pptx路径, 输出目录, "ppt")
        if ppt路径 is None:
            self.fail("LibreOffice 未生成 .ppt 夹具：私有档案 + 3 次重试后仍无产物")
        with _实现走私有档案(self.临时目录 / "实现档案"):
            结果 = 解析演示文稿(str(ppt路径))
        self.assertTrue(结果.成功, 结果.错误说明)
        文档 = 结果.值
        self.assertEqual(文档.文档类型, "演示文稿")
        self.assertEqual(文档.格式, "ppt")
        self.assertEqual(文档.解析方式, "转换")
        self.assertGreaterEqual(len(文档.块列表), 3)
        全部文本 = "\n".join(块.文本 for 块 in 文档.块列表)
        self.assertIn("中文演示文稿测试", 全部文本)
        self.assertIn("python-pptx 中文适配", 全部文本)
        self.assertIn("python-pptx", 文档.提供者版本)


if __name__ == "__main__":
    unittest.main()
